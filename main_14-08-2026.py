import os
import time
import re
import ast
import io
import json
import base64
import uuid
import threading
import asyncio
from datetime import datetime
from typing import TypedDict, Dict, Any, Optional, List
from langgraph.graph import StateGraph, END
from langchain_groq import ChatGroq
from langchain_core.messages import SystemMessage, HumanMessage
import pyodbc
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.responses import FileResponse
from pydantic import BaseModel
import uvicorn
from fastapi.middleware.cors import CORSMiddleware
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import tempfile
from langgraph.checkpoint.memory import MemorySaver

import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage
from reportlab.lib.units import inch

from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter

import requests
import edge_tts

checkpointer = MemorySaver()

# ==========================================
# 1. CONFIGURATION & ENVIRONMENT
# ==========================================
os.environ["GROQ_API_KEY"] = os.environ.get("GROQ_API_KEY", "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")

CONNECTION_STRING = (
    "Driver={ODBC Driver 17 for SQL Server};"
    "Server=160.25.62.109,1433;"
    "Database=SOUL30;"
    "UID=sa;"
    "PWD=msspl@123;"
    "Encrypt=no;"
    "TrustServerCertificate=yes;"
    "Connection Timeout=30;"
)

MAX_ROWS = 100
MAX_RETRIES = 5
MAX_HISTORY_TURNS = 6
RESOLVER_CONTEXT_TURNS = 4
_conn = None

STT_BASE_URL = "https://stt-server.mssplonline.in"
STT_TRANSCRIBE_FILE_URL = f"{STT_BASE_URL}/transcribe"
STT_TRANSCRIBE_PCM_URL = f"{STT_BASE_URL}/transcribe-pcm"
STT_HEALTH_URL = f"{STT_BASE_URL}/v1/stt/health"
STT_LANGUAGE = None

TTS_VOICE = "hi-IN-MadhurNeural"

# ==========================================
# REPORT STORE
# ==========================================
REPORT_STORE: Dict[str, Dict[str, Any]] = {}
REPORT_TTL_SECONDS = 600
REPORT_LOCK = threading.Lock()


def cleanup_expired_reports():
    now = time.time()
    expired = [
        rid for rid, rdata in REPORT_STORE.items()
        if now - rdata["created_at"] > REPORT_TTL_SECONDS
    ]
    for rid in expired:
        REPORT_STORE.pop(rid, None)
    if expired:
        print(f"[REPORT_STORE] Cleaned up {len(expired)} expired reports.")


def store_report(data: List[Dict], question: str, sql: str) -> str:
    cleanup_expired_reports()
    report_id = str(uuid.uuid4())
    with REPORT_LOCK:
        REPORT_STORE[report_id] = {
            "data": data,
            "question": question,
            "sql": sql,
            "created_at": time.time(),
        }
    print(f"[REPORT_STORE] Stored report_id={report_id} with {len(data)} rows.")
    return report_id


def get_report(report_id: str) -> Optional[Dict[str, Any]]:
    with REPORT_LOCK:
        rdata = REPORT_STORE.get(report_id)
        if not rdata:
            return None
        if time.time() - rdata["created_at"] > REPORT_TTL_SECONDS:
            REPORT_STORE.pop(report_id, None)
            return None
        return rdata


# ==========================================
# 2. DATABASE CONNECTION & QUERY EXECUTION
# ==========================================
BACKUP_TABLE_PATTERNS = [
    r"_backup", r"_temp", r"weedout", r"before_restore", r"_bak\b",
    r"WO_\d+_Backup", r"NewWeedout", r"WeedoutBooks", r"_copy$"
]

ALLOWED_TEMP_TABLES = {"m_member_temp"}

BINARY_COLUMNS = {"member_photo", "member_sign", "member_receipt"}

def is_backup_table(name: str) -> bool:
    lname = name.lower()
    if lname in {t.lower() for t in ALLOWED_TEMP_TABLES}:
        return False
    for pat in BACKUP_TABLE_PATTERNS:
        if re.search(pat, lname):
            return True
    return False

def create_connection():
    return pyodbc.connect(CONNECTION_STRING)

def get_connection():
    global _conn
    try:
        if _conn is None:
            print("Creating new DB connection...")
            _conn = create_connection()
        else:
            cursor = _conn.cursor()
            cursor.execute("SELECT 1")
    except Exception:
        print("Reconnecting DB...")
        try:
            if _conn: _conn.close()
        except: pass
        _conn = create_connection()
    return _conn

def validate_query(query: str):
    if not query or not query.strip():
        raise ValueError("Empty SQL")

    forbidden = ["DROP", "DELETE", "UPDATE", "INSERT", "ALTER",
                 "TRUNCATE", "MERGE", "GRANT", "REVOKE", "EXEC", "EXECUTE"]
    for word in forbidden:
        if re.search(rf"\b{word}\b", query, re.IGNORECASE):
            raise ValueError(f"Unsafe query detected: {word}")

    stripped = query.strip().lower()
    if not (stripped.startswith("select") or stripped.startswith("with")):
        raise ValueError("Only SELECT/WITH allowed")

    try:
        table_names = extract_table_names(query)
    except Exception:
        table_names = []

    for t in table_names:
        for pat in BACKUP_TABLE_PATTERNS:
            if re.search(pat, t, re.IGNORECASE):
                if t.lower() not in {a.lower() for a in ALLOWED_TEMP_TABLES}:
                    raise ValueError(
                        f"Query references backup/temp table: '{t}'. "
                        f"Only live tables or allowed temp tables are permitted."
                    )


def _convert_value(val, col_name: str = ""):
    """Convert a database value to a JSON-serialisable Python type.
    Binary/image columns are converted to base64 data URIs."""
    if val is None:
        return None
    if isinstance(val, bytes):
        b64 = base64.b64encode(val).decode('utf-8')
        return f"data:image/jpeg;base64,{b64}"
    if isinstance(val, datetime):
        return val.strftime('%Y-%m-%d %H:%M:%S')
    return val


def execute_query(query, limit=True):
    global _conn
    for attempt in range(MAX_RETRIES):
        cursor = None
        try:
            conn = get_connection()
            cursor = conn.cursor()
            cursor.execute("SELECT 1")
            time.sleep(0.1)
            cursor.execute(query)

            if not cursor.description:
                return []

            columns = [col[0] for col in cursor.description]
            rows = cursor.fetchall()
            results = []
            for row in rows:
                row_dict = {}
                for i, col in enumerate(columns):
                    row_dict[col] = _convert_value(row[i], col)
                results.append(row_dict)

            if limit and MAX_ROWS and len(results) > MAX_ROWS:
                results = results[:MAX_ROWS]
            return results

        except Exception as e:
            error_msg = str(e)
            print(f"DB Error (attempt {attempt+1}): {error_msg}")
            if "08S01" in error_msg or "Communication link failure" in error_msg or "HY000" in error_msg:
                try:
                    if _conn: _conn.close()
                except: pass
                _conn = None
                time.sleep(1)
                continue
            raise e
        finally:
            try:
                if cursor: cursor.close()
            except: pass
    raise Exception("Database failed after retries")

def run_sql(query, limit=True):
    try:
        validate_query(query)
        results = execute_query(query, limit=limit)
        return {"success": True, "rows": len(results), "data": results}
    except Exception as e:
        return {"success": False, "error": str(e)}


# ==========================================
# 2b. LIVE SCHEMA LOOKUP
# ==========================================
SCHEMA_CACHE: Dict[str, List[str]] = {}
SCHEMA_CACHE_LOCK = threading.Lock()

_IDENTIFIER_RE = re.compile(r'^[A-Za-z_][A-Za-z0-9_]*$')

def extract_table_names(sql: str) -> List[str]:
    raw = re.findall(r'\b(?:FROM|JOIN)\s+([A-Za-z_][A-Za-z0-9_\.]*)', sql, re.IGNORECASE)
    names = []
    for r in raw:
        name = r.split('.')[-1]
        if _IDENTIFIER_RE.match(name) and name not in names:
            names.append(name)
    return names

def get_table_columns(table_name: str) -> List[str]:
    if not _IDENTIFIER_RE.match(table_name):
        return []
    with SCHEMA_CACHE_LOCK:
        cached = SCHEMA_CACHE.get(table_name)
    if cached is not None:
        return cached
    try:
        query = (
            "SELECT COLUMN_NAME FROM INFORMATION_SCHEMA.COLUMNS "
            f"WHERE TABLE_NAME = '{table_name}' ORDER BY ORDINAL_POSITION"
        )
        results = execute_query(query, limit=False)
        cols = [r["COLUMN_NAME"] for r in results]
        with SCHEMA_CACHE_LOCK:
            SCHEMA_CACHE[table_name] = cols
        return cols
    except Exception as e:
        print(f"[SCHEMA] Could not fetch columns for {table_name}: {e}")
        return []

def find_similar_table_names(fragment: str) -> List[str]:
    safe_fragment = re.sub(r'[^A-Za-z0-9_]', '', fragment)[:60]
    if not safe_fragment:
        return []
    try:
        query = (
            "SELECT TABLE_NAME FROM INFORMATION_SCHEMA.TABLES "
            f"WHERE TABLE_NAME LIKE '%{safe_fragment}%'"
        )
        results = execute_query(query, limit=False)
        return [r["TABLE_NAME"] for r in results][:10]
    except Exception as e:
        print(f"[SCHEMA] Could not search similar table names: {e}")
        return []

def build_schema_error_hint(sql: str, error_msg: str) -> str:
    hints = []
    lower_err = error_msg.lower()

    if "invalid column name" in lower_err:
        for t in extract_table_names(sql):
            cols = get_table_columns(t)
            if cols:
                hints.append(f"{t}: {', '.join(cols)}")

    obj_match = re.search(r"Invalid object name '([^']+)'", error_msg, re.IGNORECASE)
    if obj_match:
        bad_name = obj_match.group(1).split('.')[-1]
        matches = find_similar_table_names(bad_name)
        if matches:
            hints.append(f"No table/view called '{bad_name}'. Closest real names: {', '.join(matches)}")

    return "\n".join(hints)


# ==========================================
# 3. THE SUPER-OPTIMIZED SYSTEM PROMPT
# ==========================================
SYSTEM_PROMPT = """You are an expert T-SQL Developer for SOUL 3.0 Library Management System. Convert natural language questions into accurate T-SQL queries.

==============================
SECTION 1: ARCHITECTURE — MARC 21 FORMAT
==============================
Book metadata (Title, Author, etc.) is stored as ROWS in `Biblidetails`, NOT columns. Use multiple self-LEFT JOINs to retrieve multiple fields.

MARC TAG DICTIONARY (Exact for this DB):
- 245 / a  → Title (Book/Journal Name)
- 100 / a  → Author
- 260 / a  → Publication Place
- 260 / b  → Publisher
- 260 / c  → Publication Year
- 250 / a  → Edition
- 300 / a  → Pages
- 022 / a  → ISSN
- 082 / a  → DDC / Classification Number
- 653 / a  → Subject
- 310 / a  → Frequency (e.g., Quarterly, Monthly)
- 210 / a  → Short Title
- 222 / a  → Main Title

==============================
SECTION 2: CRITICAL DATA TYPES & JOIN RULES (MUST FOLLOW)
==============================
- `t_issue.acc_no`, `t_receive.accn_no`, `t_book_transfer.acc_no`, `t_replace.old_accco` are **char** with TRAILING SPACES (e.g., '121613     ').
- `Location.p852` is **nvarchar** WITHOUT trailing spaces.
- `m_member.mem_cd` is `char` WITH TRAILING SPACES.

MANDATORY JOIN RULE: ALWAYS use `RTRIM()` when joining ANY of these char columns.
  CORRECT: `JOIN Location L ON L.p852 = RTRIM(t.acc_no)`
  WRONG:   `JOIN Location L ON L.p852 = t.acc_no`

==============================
SECTION 3: EXACT STATUS / ENUM CODES (CRITICAL FOR FILTERS)
==============================
- `Location.Status` = 'AV' (Available). Use `WHERE L.Status = 'AV'` for available books.
- `m_member.mem_status` = 'A' (Active).
- `t_book_transfer.Status` = 'T' (Transferred), 'R' (Returned)
- `t_receive.recv_status` = EMPTY STRING ' ' or NULL.
  CRITICAL: DO NOT use `WHERE recv_status = 'R'`.
  To check if a book is returned, check if it EXISTS in `t_receive`:
  `EXISTS (SELECT 1 FROM t_receive r WHERE r.accn_no = t.acc_no AND r.mem_cd = t.mem_cd AND r.iss_dt = t.iss_dt)`

==============================
SECTION 4: PREFERRED SHORTCUT VIEWS
==============================
USE these views for complex queries instead of writing 5+ JOINs:
1. v_BookDetails — AccNo, ClassNo, Title, Author, Location, LocId, Department, Status
2. VMember — mem_cd, mem_firstnm, mem_lstnm, mem_email, mem_prmntphone, mem_dept, Fclty_dept_dscr (Dept Name), mem_status, mem_ctgry, ctgry_desc, Branch_name
3. VCmpltIssueDetails — mem_cd, acc_no, iss_dt, due_dt, FValue (Title), mem_firstnm, mem_lstnm, mem_dept
4. vlocation — recID, p852, status, lastoperateddt

==============================
SECTION 5: SCHEMA MAP & EXACT RELATIONSHIPS (VERIFIED)
==============================
**m_member**: mem_cd, mem_firstnm, mem_lstnm, mem_status, mem_dept, mem_email
**t_issue**: mem_cd, acc_no, iss_dt, due_dt
**t_receive**: mem_cd, accn_no, iss_dt, recv_dt
**t_reserve**: mem_cd, record_no, resv_dt, hold_dt
**t_bookbankissue**: mem_cd, acc_no, issue_dt, due_dt
**t_bookbankreturn**: mem_cd, acc_no, issue_date, return_date
**t_memfine**: mem_cd, accn_no, fine_amt, slip_dt, fine_desc
**t_replace**: old_accno, mem_cd, title, author
**Location**: RecID, p852 (Accession No), Status ('AV'), DateofAcq
**Biblidetails**: RecID, Tag ('245'=Title, '100'=Author, '653'=Subject), SbFld ('a'), FValue

==============================
SECTION 5B: m_member_temp — EXTENDED MEMBER DETAILS TABLE
==============================
The table `m_member_temp` contains FULL member profile data. Use this table whenever the user asks for
phone, date of birth, signature, photo, address, gender, or any extended member detail.

**m_member_temp columns:**
- recordid (int, auto-increment primary key)
- mem_firstnm (first name)
- mem_lstnm (last name)
- mem_ctgry (member category code)
- mem_inst (institute)
- mem_dept (department code)
- mem_degree (degree)
- mem_year (year)
- mem_status (status: 'A'=Active, 'IA'=Inactive)
- mem_prmntadd1, mem_prmntadd2 (permanent address lines)
- mem_prmntcity (permanent city)
- mem_prmntpin (permanent PIN code)
- mem_prmntphone (permanent phone number)
- mem_tmpadd1, mem_tmpadd2 (temporary address lines)
- mem_tmpcity (temporary city)
- mem_tmppin (temporary PIN code)
- mem_tmpphone (temporary phone number)
- mem_email (email address)
- mem_id (member ID)
- remarks
- date_of_birth (datetime)
- mem_gender (Male/Female)
- mem_type (member type, e.g., 'research scholar', 'Faculty')
- member_photo (BINARY IMAGE — JPEG)
- member_sign (BINARY IMAGE — JPEG — member's signature)
- member_receipt (BINARY IMAGE — fee receipt)
- mem_photo_path (file path for photo)
- RegDate (registration date)
- feercptno (fee receipt number)
- feercptdate (fee receipt date)
- UpdatedDate (last updated date)
- mem_hosteladress (hostel address)
- mem_hostelroomno (hostel room number)

FIELD MAPPING (when user asks for ... → select this column):
  - Phone number       → mem_prmntphone (permanent) or mem_tmpphone (temporary)
  - Date of birth / DOB → date_of_birth
  - Signature           → member_sign (BINARY IMAGE)
  - Photo               → member_photo (BINARY IMAGE)
  - Email               → mem_email
  - Address (permanent) → mem_prmntadd1, mem_prmntadd2, mem_prmntcity, mem_prmntpin
  - Address (temporary) → mem_tmpadd1, mem_tmpadd2, mem_tmpcity, mem_tmppin
  - Gender              → mem_gender
  - Hostel address      → mem_hosteladress, mem_hostelroomno
  - Category            → mem_ctgry
  - Degree              → mem_degree
  - Member type         → mem_type
  - Fee receipt no      → feercptno

CRITICAL RULES FOR BINARY COLUMNS:
  1. member_sign, member_photo, member_receipt are BINARY image columns.
  2. ONLY SELECT them — NEVER use them in WHERE, JOIN, GROUP BY, or ORDER BY clauses.
  3. When the user asks for a signature or photo, include the member's name alongside the binary column
     so the response can be properly attributed.
     CORRECT: SELECT mem_firstnm, mem_lstnm, member_sign FROM m_member_temp WHERE mem_firstnm LIKE '%NAME%'
  4. Do NOT SELECT * from m_member_temp unless explicitly asked — it returns huge binary blobs.
     Always select only the specific columns the user asked for.

SEARCHING m_member_temp BY NAME:
  - If user provides ONLY FIRST NAME:
    WHERE mem_firstnm LIKE '%UNNATI%' OR mem_lstnm LIKE '%UNNATI%'
  - If user provides FULL NAME (e.g., "UNNATI SINGH"):
    WHERE mem_firstnm LIKE '%UNNATI%' AND mem_lstnm LIKE '%SINGH%'

JOINING m_member_temp WITH m_member:
  JOIN m_member_temp T ON T.mem_firstnm = M.mem_firstnm AND T.mem_lstnm = M.mem_lstnm
  (or use RTRIM on both sides if there are trailing spaces)

==============================
SECTION 6: STRICT RULES
==============================
1. Output ONLY valid T-SQL. No markdown fences, no comments, no explanations.
2. ALWAYS use RTRIM() for char↔nvarchar joins.
3. Start book-search queries FROM Biblidetails and LEFT JOIN Location.
4. Use LIKE '%...%' for text search. NEVER use exact =.
5. Use COUNT(DISTINCT RecID) for counting books.
6. ALWAYS alias selected columns uniquely with AS.
7. NEVER reference backup tables (_backup, weedout, _copy). m_member_temp IS ALLOWED.
8. For "Top N books", use SELECT DISTINCT or GROUP BY.
9. COUNTING ISSUES: When counting how many times a book was issued, DO NOT join Location or Biblidetails if not strictly necessary.
   CORRECT: SELECT TOP 10 acc_no, COUNT(*) FROM t_issue GROUP BY acc_no
10. MEMBER NAME SEARCH LOGIC:
    - If user provides ONLY FIRST NAME: WHERE mem_firstnm LIKE '%UNNATI%' OR mem_lstnm LIKE '%UNNATI%'
    - If user provides FULL NAME (e.g., "UNNATI SINGH"): ALWAYS use AND condition.
      WHERE mem_firstnm LIKE '%UNNATI%' AND mem_lstnm LIKE '%SINGH%'
11. TOTAL LIBRARY SIZE: If user asks for "total books in library", use SELECT COUNT(*) FROM Location.
12. AGGREGATE vs DETAIL: If user asks "how many" or "total", return a single scalar number using a subquery.
13. SINGLE QUERY RULE: NEVER write multiple SELECT statements in one response.
    ALWAYS combine them into a SINGLE query using LEFT JOIN or UNION ALL.
14. DAILY/MONTHLY TRENDS (ISSUED vs RETURNED):
    If user asks for "daily issued return count", DO NOT JOIN t_issue and t_receive directly.
    Instead, query them separately using UNION ALL and then group by date.
15. COMPLETE DAILY CIRCULATION FLOW (List of transactions):
    If user asks for "poora circulation flow" or "all transactions of a day", you MUST fetch BOTH issues and returns using UNION ALL. Add a Transaction_Type column.
16. EXACT ID LOOKUPS (CRITICAL):
    If user provides an exact Member ID or Accession Number, ALWAYS use RTRIM() or LIKE '%' to handle trailing spaces.
    CORRECT: WHERE RTRIM(mem_cd) = 'DPDCDC160001'
17. MEMBER DETAILS FROM m_member_temp:
    When user asks for phone, DOB, signature, photo, address, email, gender, or ANY extended member detail,
    query m_member_temp directly. Select ONLY the columns the user asked for (plus name for identification).
    DO NOT SELECT * — it returns huge binary image data.
    Example (phone + DOB):
      SELECT mem_firstnm AS First_Name, mem_lstnm AS Last_Name,
             mem_prmntphone AS Phone, date_of_birth AS Date_Of_Birth
      FROM m_member_temp
      WHERE mem_firstnm LIKE '%MAULI%' AND mem_lstnm LIKE '%SHREE%'
    Example (signature):
      SELECT mem_firstnm AS First_Name, mem_lstnm AS Last_Name, member_sign AS Signature_Image
      FROM m_member_temp
      WHERE mem_firstnm LIKE '%MAULI%'

==============================
SECTION 7: COMPLEX QUERY EXAMPLES
==============================

--- Example A: Find available books by title ---
SELECT
    B.FValue AS Title,
    A.FValue AS Author,
    L.p852 AS Accession_No
FROM Biblidetails B
LEFT JOIN Biblidetails A ON A.RecID = B.RecID AND A.Tag = '100' AND A.SbFld = 'a'
LEFT JOIN Location L ON L.RecID = B.RecID
WHERE B.Tag = '245' AND B.SbFld = 'a'
  AND B.FValue LIKE '%data structure%'
  AND L.Status = 'AV';

--- Example B: COMPLETE Issue/Return History of a Book ---
SELECT
    t.iss_dt AS Issue_Date,
    t.due_dt AS Due_Date,
    r.recv_dt AS Return_Date,
    M.mem_firstnm + ' ' + M.mem_lstnm AS Member_Name
FROM t_issue t
LEFT JOIN t_receive r ON RTRIM(r.accn_no) = RTRIM(t.acc_no) AND r.mem_cd = t.mem_cd AND r.iss_dt = t.iss_dt
JOIN m_member M ON M.mem_cd = RTRIM(t.mem_cd)
WHERE RTRIM(t.acc_no) = '150315'
ORDER BY t.iss_dt DESC;

--- Example C: Member phone, email, DOB, and address from m_member_temp ---
SELECT
    mem_firstnm AS First_Name,
    mem_lstnm AS Last_Name,
    mem_prmntphone AS Phone,
    mem_email AS Email,
    date_of_birth AS Date_Of_Birth,
    mem_gender AS Gender,
    mem_prmntadd1 AS Address_Line_1,
    mem_prmntcity AS City,
    mem_prmntpin AS PIN_Code
FROM m_member_temp
WHERE mem_firstnm LIKE '%MAULI%' AND mem_lstnm LIKE '%SHREE%';

--- Example D: Member signature image ---
SELECT
    mem_firstnm AS First_Name,
    mem_lstnm AS Last_Name,
    member_sign AS Signature_Image
FROM m_member_temp
WHERE mem_firstnm LIKE '%MAULI%' AND mem_lstnm LIKE '%SHREE%';

--- Example E: Top 10 Most Issued Books ---
SELECT TOP 10
    t.acc_no AS Accession_No,
    COUNT(*) AS Issue_Count
FROM t_issue t
GROUP BY t.acc_no
ORDER BY Issue_Count DESC;

--- Example F: Daily Issue and Return Count for a Month ---
SELECT
    Date,
    SUM(Issued) AS Issued_Books,
    SUM(Returned) AS Returned_Books
FROM (
    SELECT CAST(iss_dt AS DATE) AS Date, COUNT(*) AS Issued, 0 AS Returned
    FROM t_issue
    WHERE iss_dt >= '2026-05-01' AND iss_dt < '2026-06-01'
    GROUP BY CAST(iss_dt AS DATE)
    UNION ALL
    SELECT CAST(recv_dt AS DATE) AS Date, 0 AS Issued, COUNT(*) AS Returned
    FROM t_receive
    WHERE recv_dt >= '2026-05-01' AND recv_dt < '2026-06-01'
    GROUP BY CAST(recv_dt AS DATE)
) sub
GROUP BY Date
ORDER BY Date;

==============================
SECTION 8: CHART GENERATION DATA FORMAT
==============================
If the user asks for a "report", "graph", "chart", "trend", or "distribution", you MUST generate an SQL query that returns EXACTLY two columns named Label and Value.
- Label: The category or time period.
- Value: The numeric count or sum.

CRITICAL: NEVER use aliases like Department, Count, Active_Students, or Total for report queries. ALWAYS use Label and Value.

CRITICAL EXCEPTION:
If the user is just searching, listing, or asking "how many", DO NOT use Label and Value aliases. Use normal descriptive aliases like Title, Total_Students, Department, Issued_Books, Returned_Books.
"""

# ==========================================
# 3b. FOLLOW-UP RESOLVER PROMPT
# ==========================================
FOLLOWUP_RESOLVER_PROMPT = """You are a query-rewriting module for a Library Management System chatbot.
Your ONLY job is to look at the conversation history and the user's CURRENT question, and decide whether the
current question is a follow-up that depends on context from earlier turns (pronouns, "unhe", "unka", "uska",
"usme se", "iski", "unhe", "in members ka", "list them", "and their emails", "wapas", "usne", "ismein",
"pichle wale", implicit topic continuation, etc.).

RULES:
1. If there is NO conversation history, OR the current question is already fully standalone, output the question EXACTLY AS-IS.
2. If the current question DOES depend on earlier context, rewrite it into one fully standalone question by substituting in the specific entity/topic/filter from the conversation history.
3. NEVER answer the question. NEVER add SQL. ONLY output the rewritten natural-language question.
4. Preserve the original language style of the CURRENT question.
5. Do not invent facts that aren't implied by the history. Keep the rewrite concise and faithful.
6. Output ONLY the rewritten question text. No quotes, no markdown, no explanation.

CONVERSATION HISTORY (oldest to newest):
{history}

CURRENT QUESTION:
{question}

Rewritten standalone question:"""

# ==========================================
# 3c. FOLLOW-UP SUGGESTIONS PROMPT
# ==========================================
FOLLOWUP_SUGGESTIONS_PROMPT = """You are a helpful assistant for a Library Management System chatbot. Based on
the user's question, the SQL that was run, and the final answer just given, suggest 3 to 4 short, natural
FOLLOW-UP questions the user might logically want to ask next.

GUIDELINES:
1. Suggestions must be a sensible NEXT step from what was just answered.
2. If the answer was empty / "no records found", suggest a broader or corrected search.
3. Write each suggestion the way an actual library staff member or student would type it — short and casual, in the
   SAME language/style as the CURRENT QUESTION.
4. Do NOT repeat or rephrase a question already present in the CONVERSATION HISTORY below.
5. Keep each suggestion under ~15 words.
6. Output ONLY a JSON array of strings — nothing else.
["question one", "question two", "question three"]

CONVERSATION HISTORY (oldest to newest, avoid repeating these):
{history}

CURRENT QUESTION: {question}
RESOLVED/STANDALONE QUESTION USED FOR SQL: {resolved_question}

SQL EXECUTED:
{sql}

FINAL ANSWER GIVEN TO THE USER:
{answer}

JSON array of 3-4 follow-up questions:"""


# ==========================================
# 4. LANGGRAPH STATE & NODES
# ==========================================
class ChatTurn(TypedDict):
    question: str
    answer: str

class AgentState(TypedDict):
    question: str
    resolved_question: str
    sql_query: str
    previous_sql: str
    query_result: str
    report_data: Optional[List[Dict[str, Any]]]
    attempts: int
    error: str
    error_history: List[str]
    schema_hint: str
    chart_base64: Optional[str]
    suggested_followups: List[str]
    chat_history: List[ChatTurn]

llm = ChatGroq(model="llama-3.3-70b-versatile", temperature=0)


def resolve_followup_node(state: AgentState):
    print("---RESOLVING FOLLOW-UP---")
    question = state["question"]
    history: List[ChatTurn] = state.get("chat_history", []) or []

    if not history:
        return {"resolved_question": question}

    recent = history[-RESOLVER_CONTEXT_TURNS:]
    history_text = "\n".join(
        f"Q: {h['question']}\nA: {h['answer']}" for h in recent
    )

    prompt = FOLLOWUP_RESOLVER_PROMPT.format(history=history_text, question=question)

    try:
        response = llm.invoke([HumanMessage(content=prompt)])
        resolved = (response.content or "").strip().strip('"').strip("'").strip()
        if not resolved:
            resolved = question
    except Exception as e:
        print(f"[RESOLVER ERROR] Falling back to raw question: {e}")
        resolved = question

    return {"resolved_question": resolved}


def generate_sql_node(state: AgentState):
    print("---GENERATING SQL---")
    prev_err = state.get("error", "")
    prev_sql = state.get("previous_sql", "")
    err_hist = state.get("error_history", [])
    schema_hint = state.get("schema_hint", "")

    active_question = state.get("resolved_question") or state["question"]

    if prev_err:
        schema_block = (
            f"ACTUAL LIVE DATABASE SCHEMA for the table(s) you used, fetched just now from "
            f"INFORMATION_SCHEMA — this is GROUND TRUTH. Use ONLY these exact names. "
            f"Do NOT invent, assume, or guess any column or table name not listed here:\n{schema_hint}\n\n"
            if schema_hint else ""
        )
        user_msg = (
            f"User Question: {active_question}\n\n"
            f"Your PREVIOUS SQL (which FAILED):\n{prev_sql}\n\n"
            f"ERROR returned by SQL Server:\n{prev_err}\n\n"
            f"All errors so far: {err_hist}\n\n"
            f"{schema_block}"
            f"Fix the SQL using the real schema above if provided. Re-think the JOINs and column names. "
            f"Remember the RTRIM and Status='AV' rules. "
            f"Output ONLY the SQL without any markdown."
        )
    else:
        user_msg = (
            f"User Question: {active_question}\n\n"
            f"Generate the T-SQL query. Output ONLY the SQL without any markdown."
        )

    messages = [
        SystemMessage(content=SYSTEM_PROMPT),
        HumanMessage(content=user_msg),
    ]
    response = llm.invoke(messages)
    return {
        "sql_query": response.content,
        "previous_sql": response.content,
        "attempts": state.get("attempts", 0) + 1,
    }


def _make_llm_friendly_data(data: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Replace base64 image strings with short placeholders so the LLM context stays small."""
    llm_data = []
    for row in data:
        llm_row = {}
        for k, v in row.items():
            if isinstance(v, str) and v.startswith("data:image"):
                llm_row[k] = f"[IMAGE DATA AVAILABLE — {len(v)} chars of base64 encoded image]"
            else:
                llm_row[k] = v
        llm_data.append(llm_row)
    return llm_data


def execute_sql_node(state: AgentState):
    print("---EXECUTING SQL---")
    clean_sql = state["sql_query"].replace("```sql", "").replace("```", "").strip()
    print(f"\n[DEBUG] SQL (attempt {state.get('attempts', 0)}):\n{clean_sql}\n")

    result = run_sql(clean_sql)
    if result["success"]:
        raw_data = result["data"]

        llm_data = _make_llm_friendly_data(raw_data)

        data_str = str(llm_data)
        if len(data_str) > 12000:
            data_str = data_str[:12000] + f"\n... [TRUNCATED, total rows: {len(raw_data)}]"
        return {
            "query_result": data_str,
            "report_data": raw_data,
            "error": "",
            "schema_hint": "",
        }
    else:
        print(f"[DEBUG] SQL ERROR: {result['error']}\n")
        hist = state.get("error_history", [])
        hist.append(result["error"])

        schema_hint = ""
        try:
            schema_hint = build_schema_error_hint(clean_sql, result["error"])
            if schema_hint:
                print(f"[SCHEMA] Ground-truth hint for retry:\n{schema_hint}\n")
        except Exception as e:
            print(f"[SCHEMA] Hint lookup failed, continuing without it: {e}")

        return {
            "error": result["error"],
            "error_history": hist,
            "report_data": None,
            "schema_hint": schema_hint,
        }


def check_status_node(state: AgentState):
    if state.get("error"):
        print("---ERROR FOUND, RETRYING---")
        if state.get("attempts", 0) >= MAX_RETRIES:
            print("---MAX ATTEMPTS REACHED, STOPPING---")
            return END
        return "retry_generate"
    else:
        print("---SUCCESS---")
        return "generate_chart"


def generate_chart_node(state: AgentState):
    print("---CHECKING FOR CHART DATA---")
    data_str = state.get("query_result", "[]")
    chart_b64 = None

    try:
        data = ast.literal_eval(data_str)
        if len(data) > 1 and "Label" in data[0] and "Value" in data[0]:
            print("---GENERATING CHART---")
            labels = [d["Label"] for d in data]
            values = [d["Value"] for d in data]

            plt.figure(figsize=(10, 5))

            if len(data) <= 6:
                plt.pie(values, labels=labels, autopct='%1.1f%%', startangle=90)
                plt.axis('equal')
                plt.title('Library Data Distribution')
            else:
                plt.plot(labels, values, marker='o', linestyle='-', color='b')
                plt.fill_between(labels, values, color='skyblue', alpha=0.4)
                plt.title('Library Data Trend')
                plt.xlabel('Time Period / Category')
                plt.ylabel('Count / Value')
                plt.xticks(rotation=45)
                plt.grid(True, linestyle='--', alpha=0.7)
                plt.tight_layout()

            buf = io.BytesIO()
            plt.savefig(buf, format='png')
            plt.close()
            buf.seek(0)
            chart_b64 = base64.b64encode(buf.read()).decode('utf-8')

    except Exception as e:
        print(f"Chart generation skipped: {e}")

    return {"chart_base64": chart_b64}


def generate_answer_node(state: AgentState):
    print("---GENERATING FINAL ANSWER---")

    chart_note = ""
    if state.get("chart_base64"):
        chart_note = "A chart has been generated and attached for this data. Mention this in your response."

    row_count_note = ""
    data_str = state.get("query_result", "[]")
    try:
        actual_data = ast.literal_eval(data_str)
        if isinstance(actual_data, list):
            row_count = len(actual_data)
            row_count_note = (
                f"\nIMPORTANT: The SQL query returned exactly {row_count} rows. "
                f"If the user asks 'how many' or for a count, you MUST use this exact number ({row_count})."
            )
    except Exception as e:
        print(f"Could not calculate row count for LLM: {e}")

    active_question = state.get("resolved_question") or state["question"]

    messages = [
        SystemMessage(content=f"""You are a helpful, professional library assistant. Based on the user's question and the SQL query result, give a clear, highly accurate, and perfectly formatted answer.

LANGUAGE & FORMATTING RULES:
1. CRITICAL: DO NOT mention or hallucinate about downloading files, Excel copies, PDF copies, or "full report on file".
2. Reply in the EXACT SAME language the user used (Hindi, English, or Hinglish).
3. If the result contains multiple rows, ALWAYS format it as a clean Markdown Table.
4. If the user asks for member details AND issue history, present the details in a table, and if history is available, present it in a separate table.
5. If the result is a count or aggregate, write a clear, concise sentence.
6. If the result is empty (e.g., []), politely say "No records found matching your query." Do not invent data.
7. For dates, format as DD-MM-YYYY.
8. For money, prefix with Rs. or ₹.
9. Do not include raw JSON or Python dictionary syntax in the output. Parse it and present it beautifully.
10. TABLE SERIAL NUMBERING: Whenever you format the result as a Markdown table with multiple rows, ALWAYS add a first column named "S.No." starting from 1.

IMAGE / SIGNATURE / PHOTO DATA HANDLING:
- If you see "[IMAGE DATA AVAILABLE — ... chars of base64 encoded image]" in the SQL Result Data, this means
  a binary image (such as a member's signature or photo) was successfully retrieved.
- In your response, MENTION that the image is available and attached in the response data.
- Example: "MAULI SHREE ki signature image response me available hai."
- Example: "Member ki photo image data ke roop me available hai."
- DO NOT try to display or describe the image content. Just confirm it is available.
- The frontend will handle displaying the actual image from the base64 data.

DATE OF BIRTH FORMATTING:
- If date_of_birth is in the result, format it as DD-MM-YYYY (e.g., 08-11-1994).

PHONE NUMBER FORMATTING:
- Present phone numbers clearly, e.g., "Phone: 7408049743"

MEMBER DETAILS TABLE:
When presenting member details from m_member_temp, use this format:
| S.No. | Name | Phone | Email | Date of Birth | Gender | Address |
Only include columns that the user asked for. Do not add extra columns.
{chart_note}"""),
        HumanMessage(content=f"User Question (as originally typed): {state['question']}\n\n"
                              f"Resolved standalone question used for SQL: {active_question}\n\n"
                              f"Executed SQL:\n{state.get('sql_query','')}\n\n"
                              f"SQL Result Data:\n{state['query_result']}\n\n"
                              f"{row_count_note}\n\n"
                              f"Format this into a readable response.")
    ]
    response = llm.invoke(messages)
    return {"query_result": response.content}


def suggest_followups_node(state: AgentState):
    print("---SUGGESTING FOLLOW-UP QUESTIONS---")
    history: List[ChatTurn] = state.get("chat_history", []) or []
    recent = history[-RESOLVER_CONTEXT_TURNS:]
    history_text = "\n".join(f"Q: {h['question']}" for h in recent) or "(none yet)"

    prompt = FOLLOWUP_SUGGESTIONS_PROMPT.format(
        history=history_text,
        question=state["question"],
        resolved_question=state.get("resolved_question") or state["question"],
        sql=state.get("sql_query", "") or "(no SQL — the previous step failed)",
        answer=state.get("query_result", ""),
    )

    suggestions: List[str] = []
    raw = ""
    try:
        response = llm.invoke([HumanMessage(content=prompt)])
        raw = (response.content or "").strip()
        raw = raw.replace("```json", "").replace("```", "").strip()
        parsed = json.loads(raw)
        if isinstance(parsed, list):
            suggestions = [str(s).strip() for s in parsed if str(s).strip()][:4]
    except Exception as e:
        print(f"[SUGGEST] JSON parse failed ({e}), trying loose fallback...")
        try:
            parsed = ast.literal_eval(raw)
            if isinstance(parsed, list):
                suggestions = [str(s).strip() for s in parsed if str(s).strip()][:4]
        except Exception as e2:
            print(f"[SUGGEST] Fallback also failed, skipping suggestions: {e2}")
            suggestions = []

    print(f"[SUGGEST] {suggestions}")
    return {"suggested_followups": suggestions}


def update_memory_node(state: AgentState):
    print("---UPDATING CONVERSATION MEMORY---")
    history: List[ChatTurn] = state.get("chat_history", []) or []
    history = history + [{
        "question": state["question"],
        "answer": state.get("query_result", ""),
    }]
    if len(history) > MAX_HISTORY_TURNS:
        history = history[-MAX_HISTORY_TURNS:]
    return {"chat_history": history}


# Compile Graph
workflow = StateGraph(AgentState)
workflow.add_node("resolve_followup", resolve_followup_node)
workflow.add_node("generate_sql", generate_sql_node)
workflow.add_node("execute_sql", execute_sql_node)
workflow.add_node("generate_chart", generate_chart_node)
workflow.add_node("generate_answer", generate_answer_node)
workflow.add_node("suggest_followups", suggest_followups_node)
workflow.add_node("update_memory", update_memory_node)

workflow.set_entry_point("resolve_followup")
workflow.add_edge("resolve_followup", "generate_sql")
workflow.add_edge("generate_sql", "execute_sql")
workflow.add_conditional_edges(
    "execute_sql",
    check_status_node,
    {
        "retry_generate": "generate_sql",
        END: END,
        "generate_chart": "generate_chart"
    }
)
workflow.add_edge("generate_chart", "generate_answer")
workflow.add_edge("generate_answer", "suggest_followups")
workflow.add_edge("suggest_followups", "update_memory")
workflow.add_edge("update_memory", END)

library_agent = workflow.compile(checkpointer=checkpointer)


# ==========================================
# 5. FASTAPI APP SETUP
# ==========================================
app = FastAPI(
    title="SOUL30 Library Text-to-SQL API",
    version="3.6"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class QueryRequest(BaseModel):
    question: str
    thread_id: str


class QueryResponse(BaseModel):
    question: str
    resolved_question: str
    sql_query: str
    answer: str
    chart_base64: Optional[str] = None
    audio_base64: Optional[str] = None
    image_data: Optional[List[Dict[str, Any]]] = None
    attempts: int
    debug_error: Optional[str] = None
    report_available: bool = False
    report_id: Optional[str] = None
    suggested_followups: List[str] = []


@app.get("/welcome")
def welcome_api():
    return {
        "message": "Welcome to the SOUL 3.0 Library AI Assistant! Here are 10 complex questions you can ask to test the system:",
        "sample_questions": [
            "MAULI SHREE naam ke member ka phone number, email aur date of birth batao.",
            "MAULI SHREE ki signature do.",
            "Anju Lata naam ki member ki poori details (phone, address, DOB, gender) dikhao.",
            "Library me total kitni physical books hain?",
            "Computer science subject par 5 available books ke title aur author dikhao.",
            "UNNATI SINGH naam ki member ki poori details do aur usne kaun si book issue ki hai uski history bhi dikhao.",
            "Computer Centre department me kitne active students hain?",
            "MUMTAJ BANO MIYA ki photo do.",
            "Pichle 1 saal me sabse zyada kitni baar issue ki gayi 3 books kaun si hain?",
            "Aise 10 members dikhao jinki books overdue hain. Member ka naam, book title aur due date do."
        ]
    }

@app.get("/")
def read_root():
    return {"status": "SOUL30 Library Text-to-SQL API v3.6 (member details + signature support) is running perfectly!"}


# ==========================================
# TTS HELPER FUNCTION
# ==========================================
def generate_audio_base64(text: str) -> Optional[str]:
    if not text:
        return None
    clean_text = re.sub(r'[*_#`]', '', text)
    try:
        async def _tts():
            communicate = edge_tts.Communicate(clean_text, TTS_VOICE)
            audio_buffer = io.BytesIO()
            async for chunk in communicate.stream():
                if chunk["type"] == "audio":
                    audio_buffer.write(chunk["data"])
            return audio_buffer
        buffer = asyncio.run(_tts())
        buffer.seek(0)
        audio_b64 = base64.b64encode(buffer.read()).decode('utf-8')
        print("[TTS] Audio generated successfully.")
        return audio_b64
    except Exception as e:
        print(f"[TTS ERROR] Failed to generate audio: {e}")
        return None


# ==========================================
# STT HELPER FUNCTION
# ==========================================
def transcribe_audio_file(audio_bytes: bytes, filename: str, content_type: str) -> str:
    files = {"file": (filename, audio_bytes, content_type)}
    data = {}
    if STT_LANGUAGE:
        data["language"] = STT_LANGUAGE

    print(f"[STT] Sending audio to {STT_TRANSCRIBE_FILE_URL} ...")
    try:
        response = requests.post(STT_TRANSCRIBE_FILE_URL, files=files, data=data, timeout=120)
    except requests.exceptions.RequestException as e:
        print(f"[STT ERROR] Request failed: {e}")
        raise HTTPException(status_code=503, detail=f"Speech-to-Text server unreachable: {e}")

    if response.status_code == 400:
        raise HTTPException(status_code=400, detail="Audio file is empty, too short, or invalid. Please try recording again.")
    if response.status_code == 429:
        raise HTTPException(status_code=503, detail="Speech-to-Text server is busy right now. Please try again in a moment.")
    if response.status_code == 500:
        raise HTTPException(status_code=502, detail="Speech-to-Text transcription failed on the server. Please try again.")

    try:
        response.raise_for_status()
    except requests.exceptions.HTTPError as e:
        print(f"[STT ERROR] HTTP error: {e}")
        raise HTTPException(status_code=502, detail=f"Speech-to-Text server error: {e}")

    payload = response.json()
    if not payload.get("success", False):
        msg = payload.get("message", "Unknown STT error")
        print(f"[STT ERROR] success=false: {msg}")
        raise HTTPException(status_code=502, detail=f"Speech-to-Text failed: {msg}")

    result = payload.get("data", {}) or {}
    transcript = (result.get("transcript") or "").strip()
    detected_lang = result.get("language")
    confidence = result.get("confidence")
    print(f"[STT] Transcript: '{transcript}' | language={detected_lang} | confidence={confidence}")

    if not transcript:
        raise HTTPException(status_code=400, detail="No speech detected in the audio. Please try speaking again.")

    return transcript


# ==========================================
# IMAGE DATA EXTRACTOR
# ==========================================
def extract_image_data(report_data: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Walk through query results and pull out any base64 image fields."""
    images = []
    if not report_data:
        return images
    for idx, row in enumerate(report_data):
        first_name = row.get("First_Name") or row.get("mem_firstnm") or ""
        last_name = row.get("Last_Name") or row.get("mem_lstnm") or ""
        member_name = f"{first_name} {last_name}".strip() or f"Row {idx+1}"
        for col, val in row.items():
            if isinstance(val, str) and val.startswith("data:image"):
                label_map = {
                    "member_sign": "Signature",
                    "member_photo": "Photo",
                    "member_receipt": "Fee Receipt",
                    "Signature_Image": "Signature",
                    "Photo_Image": "Photo",
                }
                label = label_map.get(col, col.replace("_", " ").title())
                images.append({
                    "row_index": idx,
                    "member_name": member_name,
                    "column": col,
                    "label": label,
                    "mime_type": "image/jpeg",
                    "base64": val,
                })
    return images


# ==========================================
# PROCESS AGENT REQUEST (shared by /ask and /ask_audio)
# ==========================================
def process_agent_request(user_question: str, thread_id: str) -> QueryResponse:
    print(f"\n[API] User Question: {user_question}\n")

    final_state = library_agent.invoke(
        {
            "question": user_question,
            "resolved_question": "",
            "attempts": 0,
            "error": "",
            "previous_sql": "",
            "error_history": [],
            "schema_hint": "",
            "chart_base64": None,
            "report_data": None,
            "suggested_followups": [],
        },
        config={
            "configurable": {
                "thread_id": thread_id
            }
        }
    )

    final_answer = final_state.get("query_result", "Agent failed to generate answer.")
    executed_sql = (final_state.get("sql_query", "")
                    .replace("```sql", "").replace("```", "").strip())
    db_error = final_state.get("error", None) or None
    chart_b64 = final_state.get("chart_base64", None)
    resolved_q = final_state.get("resolved_question", "") or user_question
    suggested_followups = final_state.get("suggested_followups", []) or []

    audio_b64 = generate_audio_base64(final_answer)

    report_available = False
    report_id = None
    report_data = final_state.get("report_data", None)

    image_data_list = []
    if report_data and isinstance(report_data, list) and len(report_data) > 0:
        try:
            report_id = store_report(
                data=report_data,
                question=user_question,
                sql=executed_sql,
            )
            report_available = True
        except Exception as e:
            print(f"[REPORT_STORE] Failed to store report: {e}")
            report_available = False

        image_data_list = extract_image_data(report_data)
        if image_data_list:
            print(f"[IMAGE_DATA] Extracted {len(image_data_list)} image(s) from query results.")

    return QueryResponse(
        question=user_question,
        resolved_question=resolved_q,
        sql_query=executed_sql,
        answer=final_answer,
        chart_base64=chart_b64,
        audio_base64=audio_b64,
        image_data=image_data_list if image_data_list else None,
        attempts=final_state.get("attempts", 0),
        debug_error=db_error,
        report_available=report_available,
        report_id=report_id,
        suggested_followups=suggested_followups,
    )


@app.post("/ask", response_model=QueryResponse)
def ask_library_agent(request: QueryRequest):
    return process_agent_request(request.question, request.thread_id)


@app.post("/ask_audio", response_model=QueryResponse)
def ask_audio_library_agent(
    audio: UploadFile = File(...),
    thread_id: str = Form(...)
):
    print(f"\n[API] Received audio file: {audio.filename} for thread_id: {thread_id}")

    audio_bytes = audio.file.read()
    transcribed_text = transcribe_audio_file(
        audio_bytes=audio_bytes,
        filename=audio.filename or "recording.wav",
        content_type=audio.content_type or "application/octet-stream",
    )

    print(f"[AUDIO API] Transcription successful: {transcribed_text}")
    return process_agent_request(transcribed_text, thread_id)


@app.get("/stt/health")
def stt_health():
    try:
        resp = requests.get(STT_HEALTH_URL, timeout=10)
        resp.raise_for_status()
        return resp.json()
    except requests.exceptions.RequestException as e:
        raise HTTPException(status_code=503, detail=f"STT server health check failed: {e}")


@app.post("/reset-memory/{thread_id}")
def reset_memory(thread_id: str):
    try:
        library_agent.update_state(
            config={"configurable": {"thread_id": thread_id}},
            values={"chat_history": []},
        )
        return {"status": "ok", "detail": f"Memory cleared for thread_id={thread_id}"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Could not reset memory: {e}")


# ==========================================
# 6. REPORT DOWNLOAD ENDPOINTS
# ==========================================
@app.get("/report/{report_id}/excel")
def download_excel(report_id: str):
    rdata = get_report(report_id)
    if not rdata:
        raise HTTPException(status_code=404, detail="Report not found or expired.")

    try:
        data = rdata["data"]
        df = pd.DataFrame(data)
        df = df.where(pd.notnull(df), None)

        tmp_dir = tempfile.gettempdir()
        tmp_path = os.path.join(tmp_dir, f"report_{report_id}.xlsx")

        with pd.ExcelWriter(tmp_path, engine="openpyxl") as writer:
            # Clear base64 strings from dataframe
            for col in df.columns:
                df[col] = df[col].apply(
                    lambda v: "" if isinstance(v, str) and v.startswith("data:image") else v
                )
            
            df.to_excel(writer, index=False, sheet_name="Report")
            worksheet = writer.sheets["Report"]
            
            # Insert actual images into Excel cells
            for r_idx, row in enumerate(data):
                for c_idx, col_name in enumerate(row.keys()):
                    val = row[col_name]
                    if isinstance(val, str) and val.startswith("data:image"):
                        try:
                            b64_str = val.split(",")[1]
                            img_data = base64.b64decode(b64_str)
                            img_io = io.BytesIO(img_data)
                            xl_img = OpenpyxlImage(img_io)
                            xl_img.width = 120
                            xl_img.height = 120
                            cell_coord = f"{get_column_letter(c_idx + 1)}{r_idx + 2}"
                            worksheet.add_image(xl_img, cell_coord)
                            worksheet.row_dimensions[r_idx + 2].height = 90
                        except Exception as e:
                            print(f"Excel image insert error: {e}")

            for col_cells in worksheet.columns:
                max_length = max(
                    (len(str(cell.value)) if cell.value is not None else 0)
                    for cell in col_cells
                )
                col_letter = col_cells[0].column_letter
                worksheet.column_dimensions[col_letter].width = min(max_length + 2, 50)

        return FileResponse(
            path=tmp_path,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=f"library_report_{report_id[:8]}.xlsx",
        )
    except Exception as e:
        print(f"[EXCEL ERROR] {e}")
        raise HTTPException(status_code=500, detail=f"Excel generation failed: {e}")


@app.get("/report/{report_id}/pdf")
def download_pdf(report_id: str):
    rdata = get_report(report_id)
    if not rdata:
        raise HTTPException(status_code=404, detail="Report not found or expired.")

    try:
        data = rdata["data"]
        question = rdata.get("question", "")
        df = pd.DataFrame(data)

        tmp_dir = tempfile.gettempdir()
        tmp_path = os.path.join(tmp_dir, f"report_{report_id}.pdf")

        doc = SimpleDocTemplate(
            tmp_path,
            pagesize=landscape(A4),
            rightMargin=30, leftMargin=30,
            topMargin=30, bottomMargin=30,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "CustomTitle",
            parent=styles["Title"],
            fontSize=16,
            textColor=colors.HexColor("#1a5276"),
            spaceAfter=6,
        )
        sub_style = ParagraphStyle(
            "CustomSub",
            parent=styles["Normal"],
            fontSize=9,
            textColor=colors.grey,
            spaceAfter=12,
        )

        elements = []
        elements.append(Paragraph("SOUL 3.0 Library Report", title_style))
        elements.append(Paragraph(f"<b>Question:</b> {question}", sub_style))
        elements.append(Paragraph(
            f"<b>Generated:</b> {time.strftime('%d-%m-%Y %H:%M:%S')}  |  "
            f"<b>Rows:</b> {len(df)}",
            sub_style,
        ))
        elements.append(Spacer(1, 10))

        col_headers = [str(c) for c in df.columns]
        table_data = [col_headers]
        
        for _, row in df.iterrows():
            row_vals = []
            for v in row.tolist():
                if v is None or (isinstance(v, float) and pd.isna(v)):
                    row_vals.append("")
                elif isinstance(v, str) and v.startswith("data:image"):
                    try:
                        b64_str = v.split(",")[1]
                        img_data = base64.b64decode(b64_str)
                        img_io = io.BytesIO(img_data)
                        rl_img = RLImage(img_io, width=100, height=100)
                        row_vals.append(rl_img)
                    except Exception:
                        row_vals.append("[Image Error]")
                elif isinstance(v, (pd.Timestamp,)):
                    row_vals.append(v.strftime("%d-%m-%Y %H:%M"))
                else:
                    s = str(v)
                    if len(s) > 80:
                        s = s[:77] + "..."
                    row_vals.append(s)
            table_data.append(row_vals)

        page_width = landscape(A4)[0] - 60
        n_cols = max(len(col_headers), 1)
        col_width = min(page_width / n_cols, 3 * inch)
        col_widths = [col_width] * n_cols

        table = Table(table_data, repeatRows=1, colWidths=col_widths)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a5276")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
            ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#f8f9fa")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#bdc3c7")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))

        elements.append(table)
        doc.build(elements)

        return FileResponse(
            path=tmp_path,
            media_type="application/pdf",
            filename=f"library_report_{report_id[:8]}.pdf",
        )
    except Exception as e:
        print(f"[PDF ERROR] {e}")
        raise HTTPException(status_code=500, detail=f"PDF generation failed: {e}")


@app.get("/report/{report_id}/status")
def report_status(report_id: str):
    rdata = get_report(report_id)
    if not rdata:
        return {"valid": False, "detail": "Report not found or expired."}
    elapsed = time.time() - rdata["created_at"]
    remaining = max(0, REPORT_TTL_SECONDS - elapsed)
    return {
        "valid": True,
        "rows": len(rdata["data"]),
        "question": rdata["question"],
        "remaining_seconds": int(remaining),
    }


if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=7698)